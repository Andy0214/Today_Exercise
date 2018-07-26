#! /usr/bin/env python
# -*- coding: utf-8 -*-
# author = "Andy"
# Date: 2018/7/24

from docx import Document #导入库

path = "52F4E4855F36E9CBDB9EFAD032ABD250.docx" #文件路径

document = Document(path) #读入文件
tables = document.tables #获取文件中的表格集
table = tables[0]#获取文件中的第一个表格

rows1 = table.cell(0, 0).text
print("第1行第1列：", rows1)

for i in range(1, len(table.rows)):#从表格第二行开始循环读取表格数据
    result = table.cell(i, 0).text + "" + table.cell(i, 1).text
    #cell(i,0)表示第(i+1)行第1列数据，以此类推
    print(result)

# 循环每一行列数据
for row in table.rows:
    for cell in row.cells:
        print(cell.text)

row_count = len(table.rows)
col_count = len(table.columns)
print(row_count)
print(col_count)



